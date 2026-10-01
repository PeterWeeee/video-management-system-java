import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Thiết lập lề trang A4
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(13, 110, 253) # Bootstrap Primary
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(33, 37, 41)
    return p

def add_paragraph_text(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10.5)
    r.font.italic = italic
    return p

def add_code_block(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    p_pr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    p_pr.append(shd)

def add_image_with_caption(img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(f"Hình: {caption}")
        run_cap.font.name = 'Arial'
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(108, 117, 125)

# --- TRANG BÌA / TIÊU ĐỀ BÁO CÁO ---
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sch = title_p.add_run("TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT TP. HỒ CHÍ MINH\nKHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN CÔNG NGHỆ PHẦN MỀM\n\n")
r_sch.font.name = 'Arial'
r_sch.font.size = Pt(11)
r_sch.font.bold = True

r_main = title_p.add_run("BÁO CÁO BÀI THI QUÁ TRÌNH - LẬP TRÌNH WEB\nĐỀ SỐ 04 (HK1 - 2026-2027)\n\n")
r_main.font.name = 'Arial'
r_main.font.size = Pt(16)
r_main.font.bold = True
r_main.font.color.rgb = RGBColor(13, 110, 253)

# Bảng thông tin sinh viên
info_table = doc.add_table(rows=4, cols=2)
info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
info_data = [
    ("Họ và tên thí sinh:", "Nguyễn Trí Thái"),
    ("Mã số sinh viên (MSSV):", "24110330"),
    ("Mã đề thi:", "Đề số 04"),
    ("Quy tắc đặt tên file:", "Tên class_24110330.java (Áp dụng 100% các lớp, DAO, Service, Controller)")
]
for i, (label, val) in enumerate(info_data):
    row = info_table.rows[i]
    cell_lbl, cell_val = row.cells[0], row.cells[1]
    cell_lbl.width = Inches(2.2)
    cell_val.width = Inches(4.2)
    
    p0 = cell_lbl.paragraphs[0]
    r0 = p0.add_run(label)
    r0.font.name = 'Arial'
    r0.font.bold = True
    r0.font.size = Pt(10.5)
    
    p1 = cell_val.paragraphs[0]
    r1 = p1.add_run(val)
    r1.font.name = 'Arial'
    r1.font.size = Pt(10.5)
    if i < 3:
        r1.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# --- CÂU 1 ---
add_heading_1("CÂU 1 (1.5 Điểm): Cấu trúc 03 Lớp & Thiết lập Sitemesh Decorator cho User và Admin")
add_paragraph_text("Xây dựng cấu trúc dự án theo mô hình 3 tầng: Presentation Layer (MVC), Business Layer (Services), Data Access Layer (DAO). Thiết lập Sitemesh Decorators cho 02 vai trò User và Admin. Header gồm: Trang Chủ, Sản phẩm, Đăng nhập (nút viền xanh dương ở góc phải), Đăng ký (nút viền xanh lá ở góc phải), Trang quản trị (chỉ Admin mới có). Footer hiển thị chữ trắng: Họ tên: Nguyễn Trí Thái | MSSV: 24110330 | Mã đề: 04.", bold_prefix="Yêu cầu:")

add_heading_2("1.1. Cấu trúc các thành phần mã nguồn (3 Tầng):")
add_paragraph_text("vn.iotstar.dao (IUserDao_24110330, ICategoryDao_24110330, IVideoDao_24110330) và các lớp hiện thực tương ứng trong vn.iotstar.dao.impl.", bold_prefix="• Tầng DAO (Repository):")
add_paragraph_text("vn.iotstar.service (IUserService_24110330, ICategoryService_24110330, IVideoService_24110330) và vn.iotstar.service.impl.", bold_prefix="• Tầng Service:")
add_paragraph_text("vn.iotstar.controller (HomeController_24110330, AdminHomeController_24110330, ...).", bold_prefix="• Tầng Controller:")
add_paragraph_text("sitemesh3.xml, SiteMeshFilter_24110330.java, decorators/web.jsp, decorators/admin.jsp, views/common/header.jsp, footer.jsp, admin-header.jsp, admin-footer.jsp.", bold_prefix="• Tầng Decorator & Views:")

add_heading_2("1.2. Trích đoạn mã nguồn tiêu biểu:")
add_code_block("""// Header góc phải với Đăng nhập (viền xanh dương) và Đăng ký (viền xanh lá) (header.jsp):
<ul class="navbar-nav ms-auto align-items-center">
    <c:choose>
        <c:when test="${not empty sessionScope.currentUser}">
            <li class="nav-item dropdown">
                <a class="nav-link dropdown-toggle text-light" href="#" role="button" data-bs-toggle="dropdown">
                    Xin chào, <strong>${sessionScope.currentUser.fullname}</strong>
                </a>
                <ul class="dropdown-menu dropdown-menu-end">
                    <li><a class="dropdown-item text-danger" href="${pageContext.request.contextPath}/logout">Đăng xuất</a></li>
                </ul>
            </li>
        </c:when>
        <c:otherwise>
            <!-- Nút Đăng nhập viền xanh dương -->
            <li class="nav-item me-2">
                <a class="btn btn-outline-primary btn-sm px-3" href="${pageContext.request.contextPath}/login">
                    <i class="fa-solid fa-right-to-bracket"></i> Đăng nhập
                </a>
            </li>
            <!-- Nút Đăng ký viền xanh lá -->
            <li class="nav-item">
                <a class="btn btn-outline-success btn-sm px-3" href="${pageContext.request.contextPath}/register">
                    <i class="fa-solid fa-user-plus"></i> Đăng ký
                </a>
            </li>
        </c:otherwise>
    </c:choose>
</ul>

// Footer hiển thị toàn bộ chữ trắng (footer.jsp):
<footer class="footer text-center bg-dark text-white py-4 mt-5">
    <p class="mb-0 fs-5 fw-semibold text-white">Họ tên: Nguyễn Trí Thái | MSSV: 24110330 | Mã đề: 04</p>
</footer>""")

add_heading_2("1.3. Hình ảnh kết quả chạy thực tế Câu 1:")
screenshot_dir = r"C:\Users\Will\Documents\workspace-spring-tools-for-eclipse-5.3.0.RELEASE\24110330_04\doc_screenshots"
add_image_with_caption(os.path.join(screenshot_dir, "cau1_home.png"), "Giao diện Trang Chủ với Header mới (Đăng nhập viền xanh dương, Đăng ký viền xanh lá ở góc phải), Banner căn giữa và Footer trắng")
add_image_with_caption(os.path.join(screenshot_dir, "cau1_admin_home.png"), "Giao diện Bảng Điều Khiển Admin với Header Admin và Footer Họ tên, MSSV, Mã đề chữ trắng")

# --- CÂU 2 ---
add_heading_1("CÂU 2 (1.5 Điểm): Đăng Ký Kích Hoạt OTP Qua Mail, Đăng Nhập & Đăng Xuất Session")
add_paragraph_text("Trang Đăng ký có kích hoạt bằng mã OTP gửi về Email qua SMTP Gmail. Đăng nhập và Đăng xuất có sử dụng Session. Đăng nhập với vai trò admin thành công thì vào trang chủ của admin (/admin/home), User thường vào trang chủ /home. Đảm bảo toàn vẹn dữ liệu: Chỉ khi xác nhận chính xác mã OTP thì thông tin tài khoản mới được lưu vào Database; nếu gửi email thất bại hoặc email không tồn tại thì báo lỗi ngay và tuyệt đối không lưu dữ liệu rác vào Database.", bold_prefix="Yêu cầu:")

add_heading_2("2.1. Cấu trúc các thành phần mã nguồn:")
add_paragraph_text("IUserDao_24110330, UserDaoImpl_24110330 với các hàm findById, findByEmail, insert, updateActive.", bold_prefix="• Repository (DAO):")
add_paragraph_text("IUserService_24110330, UserServiceImpl_24110330 với login(username, password), activateUser(username).", bold_prefix="• Service:")
add_paragraph_text("EmailUtil_24110330 (Jakarta Mail gửi qua NguyenTT162.4@gmail.com với App Password vjea nyhm cqed jmvf, kiểm tra định dạng email và timeout kết nối), OtpUtil_24110330 (tạo OTP ngẫu nhiên 6 số).", bold_prefix="• Tiện ích Email & OTP:")
add_paragraph_text("RegisterController_24110330 (/register), VerifyOtpController_24110330 (/verify-otp), LoginController_24110330 (/login), LogoutController_24110330 (/logout).", bold_prefix="• Controller:")
add_paragraph_text("register.jsp (form đăng ký lưu giữ lại input khi lỗi), verify-otp.jsp, login.jsp.", bold_prefix="• Views:")

add_heading_2("2.2. Trích đoạn mã nguồn tiêu biểu:")
add_code_block("""// 1. RegisterController_24110330.java - Gửi OTP trước, chỉ lưu tạm vào Session (CHƯA LƯU VÀO DATABASE):
String otp = OtpUtil_24110330.generateOtp();
boolean sent = EmailUtil_24110330.sendOtpEmail(email, otp);
if (!sent) {
    request.setAttribute("error", "Không thể gửi OTP tới email (email không tồn tại hoặc lỗi máy chủ). Chưa lưu vào CSDL!");
    request.getRequestDispatcher("/WEB-INF/views/web/register.jsp").forward(request, response);
    return;
}
// Lưu tạm thông tin tài khoản vào Session để chờ xác thực:
session.setAttribute("pendingUser", pendingUser);
session.setAttribute("otpCode", otp);
response.sendRedirect(request.getContextPath() + "/verify-otp");

// 2. VerifyOtpController_24110330.java - CHỈ KHI NHẬP ĐÚNG OTP MỚI INSERT VÀO DATABASE:
if (enteredOtp != null && enteredOtp.trim().equals(sessionOtp.trim())) {
    pendingUser.setActive(true);
    userService.insert(pendingUser); // Lưu tài khoản đã kích hoạt vào Database tại đây!
    session.removeAttribute("pendingUser");
    session.removeAttribute("otpCode");
    session.setAttribute("successMessage", "Xác thực OTP thành công! Vui lòng đăng nhập.");
    response.sendRedirect(request.getContextPath() + "/login");
} else {
    request.setAttribute("error", "Mã OTP không chính xác. Vui lòng nhập lại!");
    request.getRequestDispatcher("/WEB-INF/views/web/verify-otp.jsp").forward(request, response);
}

// 3. LoginController_24110330.java - Xử lý Session và Phân quyền điều hướng:
HttpSession session = request.getSession();
session.setAttribute("currentUser", user);

if (Boolean.TRUE.equals(user.getAdmin())) {
    response.sendRedirect(request.getContextPath() + "/admin/home");
} else {
    response.sendRedirect(request.getContextPath() + "/home");
}""")

add_heading_2("2.3. Hình ảnh kết quả chạy thực tế Câu 2:")
add_image_with_caption(os.path.join(screenshot_dir, "cau2_register.png"), "Giao diện Form Đăng ký tài khoản (hỗ trợ nhập email để nhận mã OTP)")
add_image_with_caption(os.path.join(screenshot_dir, "cau2_verify_otp.png"), "Giao diện Nhập và Xác thực mã OTP 6 số gửi từ Gmail để kích hoạt tài khoản")
add_image_with_caption(os.path.join(screenshot_dir, "cau2_login.png"), "Giao diện Đăng nhập hệ thống có kiểm tra Session và phân quyền vai trò")

# --- CÂU 3 ---
add_heading_1("CÂU 3 (2.5 Điểm): Chức Năng Quản Trị CRUD Dữ Liệu Bảng Users Phân Trang 6 User/Trang")
add_paragraph_text("Xây dựng chức năng CRUD (Tạo, Xem, Cập nhật, Xóa) cho việc quản trị dữ liệu bảng Users có phân trang đúng 6 user trên 01 trang.", bold_prefix="Yêu cầu:")

add_heading_2("3.1. Cấu trúc các thành phần mã nguồn:")
add_paragraph_text("IUserDao_24110330 & UserDaoImpl_24110330 với câu truy vấn phân trang: SELECT * FROM Users ORDER BY Username ASC OFFSET ? ROWS FETCH NEXT ? ROWS ONLY và countUsers().", bold_prefix="• Repository (DAO):")
add_paragraph_text("IUserService_24110330 & UserServiceImpl_24110330 với getPage(page, 6), getTotalPages(6), insert, update, delete.", bold_prefix="• Service:")
add_paragraph_text("AdminUserController_24110330.java ánh xạ các URL: /admin/users, /admin/user/create, /admin/user/edit, /admin/user/delete.", bold_prefix="• Controller:")
add_paragraph_text("views/admin/user-list.jsp (hiển thị bảng người dùng + phân trang), views/admin/user-form.jsp (form thêm mới và chỉnh sửa).", bold_prefix="• Views:")

add_heading_2("3.2. Trích đoạn mã nguồn tiêu biểu:")
add_code_block("""// UserDaoImpl_24110330.java - Truy vấn phân trang 6 users / trang:
@Override
public List<UserModel_24110330> findPage(int page, int pageSize) {
    int offset = (page - 1) * pageSize;
    String sql = "SELECT * FROM Users ORDER BY Username ASC OFFSET ? ROWS FETCH NEXT ? ROWS ONLY";
    // ps.setInt(1, offset); ps.setInt(2, pageSize);
}

// AdminUserController_24110330.java - Xử lý phân trang PAGE_SIZE = 6:
private static final int PAGE_SIZE = 6;
int totalUsers = userService.countUsers();
int totalPages = (int) Math.ceil((double) totalUsers / PAGE_SIZE);
List<UserModel_24110330> list = userService.getPage(page, PAGE_SIZE);""")

add_heading_2("3.3. Hình ảnh kết quả chạy thực tế Câu 3:")
add_image_with_caption(os.path.join(screenshot_dir, "cau3_admin_users_p1.png"), "Danh sách người dùng phân trang Trang 1 (gồm 6 users: admin, thainguyen, user01..user04)")
add_image_with_caption(os.path.join(screenshot_dir, "cau3_admin_users_p2.png"), "Danh sách người dùng phân trang Trang 2 (gồm 6 users: user05..user10)")
add_image_with_caption(os.path.join(screenshot_dir, "cau3_admin_user_form.png"), "Form Thêm mới / Cập nhật người dùng với đầy đủ trường dữ liệu")

# --- CÂU 4 ---
add_heading_1("CÂU 4 (1.5 Điểm): Xây Dựng Trang Chi Tiết 01 Video Theo Đúng Mẫu")
add_paragraph_text("Xây dựng trang chi tiết 01 video với bố cục: Poster bên trái; bên phải gồm Tiêu đề, Mã video, Category name, View, Share(số lượng), Like(số lượng); bên dưới là description.", bold_prefix="Yêu cầu:")

add_heading_2("4.1. Cấu trúc các thành phần mã nguồn:")
add_paragraph_text("IVideoDao_24110330 & VideoDaoImpl_24110330 phương thức findByIdWithDetails(String videoId) thực hiện JOIN với Category và đếm số lượng Share từ bảng Shares, số Like từ bảng Favorites.", bold_prefix="• Repository (DAO):")
add_paragraph_text("IVideoService_24110330 & VideoServiceImpl_24110330 phương thức getVideoDetail(String videoId).", bold_prefix="• Service:")
add_paragraph_text("VideoDetailController_24110330.java ánh xạ URL /video/detail.", bold_prefix="• Controller:")
add_paragraph_text("views/web/video-detail.jsp thiết kế khung hiển thị 2 cột trên và description bên dưới chuẩn 100% mẫu đề thi.", bold_prefix="• Views:")

add_heading_2("4.2. Trích đoạn mã nguồn tiêu biểu:")
add_code_block("""// VideoDaoImpl_24110330.java - Truy vấn chi tiết Video kèm Like và Share:
String sql = "SELECT v.*, c.Categoryname, "
           + "       (SELECT COUNT(*) FROM Shares s WHERE s.VideoId = v.VideoId) AS ShareCount, "
           + "       (SELECT COUNT(*) FROM Favorites f WHERE f.VideoId = v.VideoId) AS LikeCount "
           + "FROM Videos v LEFT JOIN Category c ON v.CategoryId = c.CategoryId "
           + "WHERE v.VideoId = ?";""")

add_heading_2("4.3. Hình ảnh kết quả chạy thực tế Câu 4:")
add_image_with_caption(os.path.join(screenshot_dir, "cau4_video_detail.png"), "Trang chi tiết video hiển thị Poster, Tiêu đề, Mã, Chuyên mục, Lượt xem, Nút Share, Like và Mô tả")

# --- CÂU 5 & CÂU 6 ---
add_heading_1("CÂU 5 (2.5 Điểm) & CÂU 6 (0.5 Điểm): Hiển Thị Video Theo Category Phân Trang 3 Video/Trang & Đếm Số Lượng Video")
add_paragraph_text("Hiển thị video theo từng Category. Tiêu đề hiển thị Category Name kèm tổng số video: Category Name (Số lượng). Lưới hiển thị 3 video/hàng. Phân trang dạng << 1 2 3 4 5 >> với đúng 3 video trên 01 trang.", bold_prefix="Yêu cầu:")

add_heading_2("5.1. Cấu trúc các thành phần mã nguồn:")
add_paragraph_text("IVideoDao_24110330 (findByCategoryIdPaged, countByCategoryId) và ICategoryDao_24110330 (findCategoryVideoCounts).", bold_prefix="• Repository (DAO):")
add_paragraph_text("IVideoService_24110330 (getVideosByCategoryPaged, countVideosByCategory, getTotalPagesByCategory).", bold_prefix="• Service:")
add_paragraph_text("VideoCategoryController_24110330.java ánh xạ URL /videos/category.", bold_prefix="• Controller:")
add_paragraph_text("views/web/videos-by-category.jsp thiết kế bảng 3 cột, tiêu đề màu vàng Category Name (Count), thanh phân trang << 1 2 ... >>.", bold_prefix="• Views:")

add_heading_2("5.2. Trích đoạn mã nguồn tiêu biểu:")
add_code_block("""// VideoCategoryController_24110330.java - Xử lý phân trang 3 video / trang và đếm số lượng:
private static final int PAGE_SIZE = 3;
int totalVideos = videoService.countVideosByCategory(categoryId); // Câu 6
int totalPages = (int) Math.ceil((double) totalVideos / PAGE_SIZE);
List<VideoModel_24110330> videoList = videoService.getVideosByCategoryPaged(categoryId, page, PAGE_SIZE); // Câu 5

// CategoryDaoImpl_24110330.java - Đếm số lượng video theo từng Category (Câu 6):
String sql = "SELECT c.CategoryId, c.Categoryname, c.Categorycode, COUNT(v.VideoId) AS VideoCount "
           + "FROM Category c LEFT JOIN Videos v ON c.CategoryId = v.CategoryId "
           + "GROUP BY c.CategoryId, c.Categoryname, c.Categorycode ORDER BY c.CategoryId ASC";""")

add_heading_2("5.3. Hình ảnh kết quả chạy thực tế Câu 5 & Câu 6:")
add_image_with_caption(os.path.join(screenshot_dir, "cau5_category_p1.png"), "Trang Video theo chuyên mục Lập trình Java (6) - Trang 1 với 3 video đầu tiên và phân trang << 1 2 >>")
add_image_with_caption(os.path.join(screenshot_dir, "cau5_category_p2.png"), "Trang Video theo chuyên mục Lập trình Java (6) - Trang 2 với 3 video tiếp theo")

# --- LƯU FILE WORD ---
target_file = r"C:\ute\WEB programming\24110330.docx"
target_file_updated = r"C:\ute\WEB programming\24110330_updated.docx"
os.makedirs(os.path.dirname(target_file), exist_ok=True)
try:
    doc.save(target_file)
    print("Da cap nhat file Word thanh cong vao:", target_file)
except PermissionError:
    doc.save(target_file_updated)
    print("File 24110330.docx dang duoc mo trong Microsoft Word nen khong the ghi de truc tiep.")
    print("Da luu ban moi nhat vao:", target_file_updated)
